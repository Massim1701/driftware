#!/usr/bin/env python3
"""Phase A, UK-Charts-Variante von tools/fetch_billboard_hits.py (siehe dort
fuer den Hintergrund zum Gesamtmuster: chart-basierte Quelle statt weiterer
Discogs-Style-Suchen, um Streaming-Aera-blinde 'have'-Sortierung auszugleichen).

Nutzt statt der Billboard-Hot-100 das kostenlose, bereits fertig gescrapte
UK-Top-100-Singles-Chartarchiv von github.com/JackDanHollister/UkTop100Scrape
(Excel-Datei, eine Kalenderwoche pro Tabellenblatt, 1952-heute) um Songs zu
finden, die es in die UK-Charts geschafft haben, aber noch nicht im
(deutschsprachigen) Dekaden-Katalog stehen.

Nutzerwunsch 26.9.: "das selbe machen wir mit den UK Charts der 80er 90er"
-- gleiches Vorgehen wie beim Billboard-Import, aber ueber eine unabhaengige,
britische Chartquelle (findet z.B. reine UK-Hits, die es nie in die US-Charts
geschafft haben).

WICHTIG: nutzt dieselbe Warteliste/denselben Notfound-Cache wie
fetch_billboard_hits.py (queue/<dekade>-erweiterung.json) -- Phase B
(process_decade_queue.py) bleibt dadurch unveraendert, und ein Song, der
sowohl in den US- als auch den UK-Charts war, wird nicht doppelt als
Kandidat angelegt (beide Skripte pruefen dieselbe Warteliste vor dem
Discogs-Suchen).

Aufruf: python3 tools/fetch_uk_chart_hits.py <decade> [--min-peak N]
(Default min-peak: 100, d.h. die komplette Top 100 zaehlt als "Hit").
Zeitbudget ueber UKCHARTS_TIME_BUDGET_SECONDS steuerbar, laeuft resumable
ueber mehrere Aufrufe (Fortschritt wird laufend in die Warteliste gespeichert).

Benoetigt 'openpyxl' (pip install openpyxl)."""
import json
import os
import re
import sys
import tempfile
import time
import urllib.error
import urllib.parse
import urllib.request

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__))) if "__file__" in dir() else "."

DECADE_YEARS = {
    "70er": range(1970, 1980),
    "80er": range(1980, 1990),
    "90er": range(1990, 2000),
    "2000er": range(2000, 2010),
    "2010er": range(2010, 2020),
    "2020er": range(2020, 2026),
}

DECADE_CATALOG = {
    "70er": "70er-music/songs.json",
    "80er": "80er-music/songs.json",
    "90er": "90er-music/songs.json",
    "2000er": "2000er-music/songs.json",
    "2010er": "2010er-music/songs.json",
    "2020er": "2020er-music/songs.json",
}

# Dieselben Wartelisten wie fetch_billboard_hits.py / process_decade_queue.py
# -- bewusst KEIN eigener Queue-Name (siehe Modul-Docstring).
DECADE_QUEUE = {
    "70er": "queue/70er-erweiterung.json",
    "80er": "queue/80er-erweiterung.json",
    "90er": "queue/90er-erweiterung.json",
    "2000er": "queue/2000er-erweiterung.json",
    "2010er": "queue/2010er-erweiterung.json",
    "2020er": "queue/2020er-erweiterung.json",
}

DECADE_NOTFOUND = {
    "70er": "queue/70er-erweiterung-notfound.json",
    "80er": "queue/80er-erweiterung-notfound.json",
    "90er": "queue/90er-erweiterung-notfound.json",
    "2000er": "queue/2000er-erweiterung-notfound.json",
    "2010er": "queue/2010er-erweiterung-notfound.json",
    "2020er": "queue/2020er-erweiterung-notfound.json",
}

UK_CHARTS_REPO_CONTENTS_API = "https://api.github.com/repos/JackDanHollister/UkTop100Scrape/contents"
# Bewusst AUSSERHALB des Repos (nicht unter ROOT), gleicher Grund wie bei
# BILLBOARD_CACHE in fetch_billboard_hits.py -- kein Katalog-Inhalt, soll nie
# versehentlich committet werden.
UK_CHARTS_CACHE = os.path.join(tempfile.gettempdir(), "driftware-uk-top100-songs-cache.xlsx")
UK_CHARTS_CACHE_MAX_AGE = 24 * 3600  # 1 Tag

USER_AGENT = "DriftwareCatalogBot/1.0 +https://driftware.online"
MIN_HAVE = 3  # wie bei fetch_billboard_hits.py: Chart-Platzierung sichert schon die
# Qualitaet, keine hohe Discogs-Sammlerzahl noetig.
TIME_BUDGET_SECONDS = int(os.environ.get("UKCHARTS_TIME_BUDGET_SECONDS", "1500"))
SAVE_EVERY = 5
START_TS = time.time()


def api_get(url, retries=5):
    headers = {"User-Agent": USER_AGENT}
    discogs_token = os.environ.get("DISCOGS_TOKEN")
    if discogs_token:
        headers["Authorization"] = "Discogs token=" + discogs_token
    for attempt in range(retries):
        req = urllib.request.Request(url, headers=headers)
        try:
            with urllib.request.urlopen(req, timeout=30) as resp:
                return json.loads(resp.read().decode("utf-8"))
        except urllib.error.HTTPError as e:
            if e.code == 429:
                wait = 3 * (attempt + 1)
                print(f"    Rate-Limit (429), warte {wait}s ...")
                time.sleep(wait)
                continue
            print(f"    HTTP-Fehler {e.code} bei {url}: {e.read()[:200]}")
            return None
        except Exception as e:
            print(f"    Fehler bei {url}: {e}")
            time.sleep(2)
    return None


def song_id(artist, title):
    return (artist or "").strip().lower(), (title or "").strip().lower()


def normalize(s):
    return "".join(ch for ch in (s or "").lower() if ch.isalnum())


def artist_key(a):
    """Normalisierter Kuenstlername fuer den Katalog-Abgleich, 'The '-Praefix
    entfernt (siehe fetch_billboard_hits.py fuer die Begruendung)."""
    a = (a or "").strip().lower()
    if a.startswith("the "):
        a = a[4:]
    return normalize(a)


def title_variants(t):
    """Volltitel UND der Teil vor '/' oder '(' -- beide normalisiert (siehe
    fetch_billboard_hits.py fuer die Begruendung, z.B. Doppel-A-Seiten)."""
    t = (t or "").strip().lower()
    primary = re.split(r"[/(]", t)[0].strip()
    return {normalize(t), normalize(primary)} - {""}


def build_catalog_index(decade_key):
    path = os.path.join(ROOT, DECADE_CATALOG[decade_key])
    with open(path, "r", encoding="utf-8") as f:
        data = json.load(f)
    index = {}
    for songs in data.values():
        for s in songs:
            index.setdefault(artist_key(s.get("a")), set()).update(title_variants(s.get("t")))
    return index


def is_in_catalog(index, artist, title):
    variants = title_variants(title)
    cat_titles = index.get(artist_key(artist))
    if not cat_titles:
        return False
    if variants & cat_titles:
        return True
    return any(
        len(v) > 3 and len(ct) > 3 and (v in ct or ct in v)
        for v in variants for ct in cat_titles
    )


def find_uk_charts_download_url():
    """Ermittelt die aktuellste 'top_100_songs_*.xlsx'-Datei im Scraper-Repo.
    Der Dateiname traegt ein Enddatum, das sich bei jedem Repo-Update aendert
    und dabei uneinheitlich formatiert ist (z.B. '..._2024.xlsx' vs
    '..._20250809.xlsx') -- ueber die groesste Dateigroesse statt
    Datums-Parsing ermittelt ist robuster: mehr Chartwochen = neuere Datei."""
    req = urllib.request.Request(UK_CHARTS_REPO_CONTENTS_API, headers={"User-Agent": USER_AGENT})
    with urllib.request.urlopen(req, timeout=30) as resp:
        entries = json.loads(resp.read().decode("utf-8"))
    candidates = [e for e in entries if re.match(r"^top_100_songs.*\.xlsx$", e.get("name", ""))]
    if not candidates:
        raise RuntimeError("Keine top_100_songs-*.xlsx im UK-Charts-Repo gefunden.")
    best = max(candidates, key=lambda e: e.get("size", 0))
    return best["download_url"]


def load_uk_chart_workbook():
    import openpyxl
    if os.path.exists(UK_CHARTS_CACHE) and (time.time() - os.path.getmtime(UK_CHARTS_CACHE)) < UK_CHARTS_CACHE_MAX_AGE:
        print("UK-Top-100-Chartarchiv aus lokalem Cache geladen.", file=sys.stderr)
    else:
        print("Lade UK-Top-100-Chartarchiv (~13 MB, einmalig)...", file=sys.stderr)
        url = find_uk_charts_download_url()
        req = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
        with urllib.request.urlopen(req, timeout=120) as resp:
            data = resp.read()
        os.makedirs(os.path.dirname(UK_CHARTS_CACHE), exist_ok=True)
        tmp = UK_CHARTS_CACHE + ".tmp"
        with open(tmp, "wb") as f:
            f.write(data)
        os.replace(tmp, UK_CHARTS_CACHE)
    return openpyxl.load_workbook(UK_CHARTS_CACHE, read_only=True, data_only=True)


WEEK_SHEET_RE = re.compile(r"^Week (\d{4})(\d{2})(\d{2})$")


def collect_decade_hits(workbook, decade_key, min_peak):
    """Bester (niedrigster) Peak je Song ueber alle Wochenblaetter, deren
    Jahr in die Dekade faellt. Spaltenlayout je Blatt: Song, Artist,
    Position, Last Week, Peak, Weeks on Chart, Week -- durchgehend gleich
    (Stichprobe 1980/1985/1990/1995/1999 geprueft, daher per Header-Index
    statt fester Spaltennummer gelesen, falls sich das doch mal aendert)."""
    years = set(DECADE_YEARS[decade_key])
    best = {}
    for sheet_name in workbook.sheetnames:
        m = WEEK_SHEET_RE.match(sheet_name)
        if not m:
            continue
        year = int(m.group(1))
        if year not in years:
            continue
        ws = workbook[sheet_name]
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            continue
        try:
            idx_song = header.index("Song")
            idx_artist = header.index("Artist")
            idx_peak = header.index("Peak")
        except ValueError:
            continue
        max_idx = max(idx_song, idx_artist, idx_peak)
        for row in rows:
            if row is None or len(row) <= max_idx:
                continue
            title = (row[idx_song] or "").strip() if row[idx_song] else ""
            artist = (row[idx_artist] or "").strip() if row[idx_artist] else ""
            try:
                peak = int(str(row[idx_peak]).strip())
            except (TypeError, ValueError):
                continue
            if not artist or not title or peak > min_peak:
                continue
            key = song_id(artist, title)
            cur = best.get(key)
            if cur is None or peak < cur["peak"]:
                best[key] = {"a": artist, "t": title, "peak": peak, "year": year}
    return best


def _pick_matching_release(results, artist, title):
    a_key = artist_key(artist)
    t_variants = title_variants(title)
    for r in results:
        have = (r.get("community") or {}).get("have", 0)
        if have < MIN_HAVE:
            continue
        title_full = r.get("title") or ""
        r_artist, r_title = (title_full.split(" - ", 1) + [""])[:2] if " - " in title_full else ("", title_full)
        r_artist_key = artist_key(r_artist)
        r_t_variants = title_variants(r_title)
        title_ok = bool(t_variants & r_t_variants) or any(
            len(v) > 3 and len(rv) > 3 and (v in rv or rv in v) for v in t_variants for rv in r_t_variants
        )
        if not title_ok:
            continue
        artist_ok = bool(a_key) and bool(r_artist_key) and (
            a_key == r_artist_key or (len(a_key) > 3 and len(r_artist_key) > 3 and (a_key in r_artist_key or r_artist_key in a_key))
        )
        if artist_ok:
            return r
    return None


def discogs_search_release(artist, title):
    """Identisch zu fetch_billboard_hits.py: zuerst Discogs' eigene
    'artist'+'track'-Suchfelder (praezises serverseitiges Matching, findet
    auch Singles beruehmter Kuenstler, die sonst hinter deren eigenen
    Best-Of-Alben verschwinden), dann 'release_title', dann Freitext-'q' als
    Fallback."""
    if artist:
        params = {"artist": artist, "track": title, "type": "release", "per_page": "20"}
        url = "https://api.discogs.com/database/search?" + urllib.parse.urlencode(params)
        data = api_get(url)
        match = _pick_matching_release((data or {}).get("results", []), artist, title)
        if match:
            return match

    params = {"release_title": title, "type": "release", "sort": "have", "sort_order": "desc", "per_page": "25"}
    url = "https://api.discogs.com/database/search?" + urllib.parse.urlencode(params)
    data = api_get(url)
    match = _pick_matching_release((data or {}).get("results", []), artist, title)
    if match:
        return match

    params = {"q": f"{artist} {title}", "type": "release", "sort": "have", "sort_order": "desc", "per_page": "20"}
    url = "https://api.discogs.com/database/search?" + urllib.parse.urlencode(params)
    data = api_get(url)
    return _pick_matching_release((data or {}).get("results", []), artist, title)


def build_candidate(hit, release):
    return {
        "id": release.get("id"),
        "a": hit["a"],
        "t": hit["t"],
        "y": release.get("year") or hit.get("year"),
        "g": ", ".join(release.get("genre") or []),
        "s": ", ".join(release.get("style") or []),
        "c": release.get("country"),
        "l": ", ".join(release.get("label") or []),
        "th": release.get("thumb") or None,
        "cv": release.get("cover_image") or None,
        "u": "https://www.discogs.com" + (release.get("uri") or ""),
        "hv": (release.get("community") or {}).get("have", 0),
        "genre_key": "Ohne",
        "src": "ukcharts",
        "uk_peak": hit["peak"],
    }


def main():
    if len(sys.argv) < 2 or sys.argv[1] not in DECADE_YEARS:
        print("Usage: fetch_uk_chart_hits.py <decade> [--min-peak N]")
        print("Dekaden:", list(DECADE_YEARS))
        sys.exit(1)
    decade_key = sys.argv[1]
    min_peak = 100
    if "--min-peak" in sys.argv:
        min_peak = int(sys.argv[sys.argv.index("--min-peak") + 1])

    workbook = load_uk_chart_workbook()
    hits = collect_decade_hits(workbook, decade_key, min_peak)
    print(f"{decade_key}: {len(hits)} eindeutige UK-Chart-Hits (peak<={min_peak}) im Dekaden-Zeitraum.")

    catalog_index = build_catalog_index(decade_key)
    print(f"{decade_key}: {sum(len(v) for v in catalog_index.values())} Titel-Varianten im Katalog indiziert.")

    queue_path = os.path.join(ROOT, DECADE_QUEUE[decade_key])
    existing_queue = []
    if os.path.exists(queue_path):
        with open(queue_path, "r", encoding="utf-8") as f:
            existing_queue = json.load(f)
    queue_index = {}
    for c in existing_queue:
        queue_index.setdefault(artist_key(c.get("a")), set()).update(title_variants(c.get("t")))

    notfound_path = os.path.join(ROOT, DECADE_NOTFOUND[decade_key])
    notfound_keys = set()
    if os.path.exists(notfound_path):
        with open(notfound_path, "r", encoding="utf-8") as f:
            notfound_keys = set(tuple(k) for k in json.load(f))

    missing = [
        hit for hit in hits.values()
        if not is_in_catalog(catalog_index, hit["a"], hit["t"])
        and not is_in_catalog(queue_index, hit["a"], hit["t"])
        and song_id(hit["a"], hit["t"]) not in notfound_keys
    ]
    skipped_notfound = len(hits) - len(missing) - sum(
        1 for hit in hits.values()
        if is_in_catalog(catalog_index, hit["a"], hit["t"]) or is_in_catalog(queue_index, hit["a"], hit["t"])
    )
    if skipped_notfound > 0:
        print(f"{decade_key}: {skipped_notfound} Songs uebersprungen (bereits als 'kein Discogs-Treffer' bekannt).")
    missing.sort(key=lambda h: h["peak"])
    print(f"{decade_key}: {len(missing)} Songs fehlen noch und werden jetzt bei Discogs gesucht.")

    def save_progress(new_candidates):
        merged = existing_queue + new_candidates
        tmp = queue_path + ".tmp"
        os.makedirs(os.path.dirname(queue_path), exist_ok=True)
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(merged, f, separators=(",", ":"), ensure_ascii=False)
        os.replace(tmp, queue_path)

    def save_notfound(keys):
        merged = sorted(notfound_keys | keys)
        tmp = notfound_path + ".tmp"
        os.makedirs(os.path.dirname(notfound_path), exist_ok=True)
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump([list(k) for k in merged], f, separators=(",", ":"), ensure_ascii=False)
        os.replace(tmp, notfound_path)

    new_candidates = []
    new_notfound = set()
    not_found = 0
    i = 0
    while i < len(missing):
        if time.time() - START_TS > TIME_BUDGET_SECONDS:
            print(f"Zeitbudget ({TIME_BUDGET_SECONDS}s) erreicht, Rest bleibt fuer den naechsten Lauf.")
            break
        hit = missing[i]
        i += 1
        release = discogs_search_release(hit["a"], hit["t"])
        if release:
            new_candidates.append(build_candidate(hit, release))
            print(f"  [{i}/{len(missing)}] + {hit['a']} - {hit['t']} (Peak #{hit['peak']})")
        else:
            not_found += 1
            new_notfound.add(song_id(hit["a"], hit["t"]))
            print(f"  [{i}/{len(missing)}] kein Discogs-Treffer: {hit['a']} - {hit['t']}")
        if len(new_candidates) and len(new_candidates) % SAVE_EVERY == 0:
            save_progress(new_candidates)
        if len(new_notfound) and len(new_notfound) % SAVE_EVERY == 0:
            save_notfound(new_notfound)
        time.sleep(1.1)

    save_progress(new_candidates)
    save_notfound(new_notfound)
    print(f"Fertig fuer diesen Lauf ({decade_key}): {len(new_candidates)} neue Kandidaten in die Warteliste, "
          f"{not_found} ohne Discogs-Treffer uebersprungen, {len(missing) - i} bleiben fuer den naechsten Lauf.")


if __name__ == "__main__":
    main()
